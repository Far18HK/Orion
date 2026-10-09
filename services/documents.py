"""Extracción de texto de documentos: PDF, DOCX y archivos de texto."""
import asyncio
import io
from pathlib import Path

MAX_DOC_CHARS = 24_000  # Lo que se le pasa al modelo; más que esto se recorta

TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".log", ".xml", ".html", ".yaml", ".yml", ".ini",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".sql",
}
SUPPORTED_EXTENSIONS = {".pdf", ".docx"} | TEXT_EXTENSIONS


class DocumentError(Exception):
    """Error amigable para el usuario."""


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("Ese PDF tiene contraseña 🔒 no puedo abrirlo.")

    parts: list[str] = []
    total = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        parts.append(text)
        total += len(text)
        if total > MAX_DOC_CHARS:  # No hace falta leer cientos de páginas de más
            break
    return "\n".join(parts)


def _read_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _extract_sync(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _read_pdf(data)
    if ext == ".docx":
        return _read_docx(data)
    if ext in TEXT_EXTENSIONS:
        return _read_text(data)
    raise DocumentError(
        "Ese formato no lo puedo leer todavía 📄 Acepto PDF, DOCX y archivos de texto "
        "(txt, md, csv, json, código...)."
    )


async def extract_text(data: bytes, filename: str) -> tuple[str, bool]:
    """Devuelve (texto, fue_recortado). Se ejecuta en un hilo porque leer PDFs consume CPU."""
    try:
        text = await asyncio.to_thread(_extract_sync, data, filename)
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(
            "No pude leer ese archivo 😕 puede estar dañado o en un formato raro."
        ) from e

    text = text.strip()
    if not text:
        raise DocumentError(
            "No encontré texto en ese archivo 🤔 si es un PDF escaneado, "
            "mándalo como foto y lo leo con visión."
        )

    truncated = len(text) > MAX_DOC_CHARS
    return text[:MAX_DOC_CHARS], truncated
